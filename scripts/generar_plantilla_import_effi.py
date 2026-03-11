#!/usr/bin/env python3
"""
generar_plantilla_import_effi.py — Etapa 4a del import CRM → Effi

Lee los contactos de EspoCRM con fuente='CRM' y enviado_a_effi=0
y genera un archivo XLSX con el formato de la plantilla de importación de Effi.

El archivo generado se guarda en: /tmp/import_clientes_effi_<fecha>.xlsx

Luego el usuario (o el script Playwright) sube ese XLSX a Effi manualmente.

Mapeos de IDs (según plantilla_importacion_cliente.xlsx):
  tipo_identificacion  → ID numérico Effi (hoja "Tipos de identificación")
  ciudad               → ID numérico Effi (tabla ciudad en espocrm DB)
  tarifa_precios       → ID numérico Effi (zeffi_tarifas_precios.id)
  tipo_de_marketing    → ID numérico Effi (zeffi_tipos_marketing.id)
  tipo_persona         → 1=Física, 2=Jurídica
  regimen_tributario   → 5=No responsable IVA (natural) / 4=Responsable IVA (jurídica)
  email_responsable    → 'equipo.origensilvestre@gmail.com' (fijo)
  sucursal             → 1 (principal, fijo)
  moneda               → 'COP' (fijo)
  permitir_venta       → 1 (fijo)

Ejecutar manualmente:
  python3 scripts/generar_plantilla_import_effi.py

  # Para procesar TODOS los contactos CRM (independiente de enviado_a_effi):
  python3 scripts/generar_plantilla_import_effi.py --todos
"""

import argparse
import sys
from datetime import date
from pathlib import Path

import mysql.connector
import openpyxl
from openpyxl import Workbook

# ─── Configuración ─────────────────────────────────────────────────────────────

DB_ESPO = dict(host='127.0.0.1', port=3306, user='osadmin', password='Epist2487.', database='espocrm')
DB_EFFI = dict(host='127.0.0.1', port=3306, user='osadmin', password='Epist2487.', database='effi_data')

OUTPUT_DIR = Path('/tmp')
EMAIL_RESPONSABLE = 'equipo.origensilvestre@gmail.com'

# Mapeos fijos
TIPO_ID_MAP = {
    'Número de Identificación Tributaria CO': 1,
    'NIT': 1,                      # alias corto por si viene migrado
    'Cédula de ciudadanía': 2,
    'Cédula de Ciudadanía': 2,     # variante mayúscula
    'Pasaporte': 3,
    'Cédula de extranjería': 4,
    'Cédula de Extranjería': 4,
}
TIPO_PERSONA_MAP = {
    'Física (natural)': 1,
    'Jurídica (moral)': 2,
}
# Tipo de cliente → ID numérico Effi (posición en lista de opciones)
TIPO_CLIENTE_MAP = {
    'Común':      1,
    'Fiel':       2,
    'Desertor':   3,
    'Mayorista':  4,
    'Importador': 5,
    'Industrial': 6,
}
# Régimen tributario por defecto según tipo de persona
REGIMEN_NATURAL  = 5   # No responsable del IVA
REGIMEN_JURIDICA = 4   # Responsable del IVA

# Columnas de la plantilla (36 columnas en orden exacto)
COLUMNAS = [
    'ID EFFI: Tipo de documento *',
    'Número de documento *',
    'Nombre | Razón social *',
    'Email',
    'Página web',
    'ID EFFI: Ciudad',
    'Código DANE Ciudad',
    'Dirección',
    'Teléfono 1',
    'Referencia teléfono 1',
    'Teléfono 2',
    'Referencia teléfono 2',
    'Celular',
    'WhatsApp',
    'Fecha de nacimiento',
    'Género',
    'Estado cívil',
    'Tipo de persona *',
    'Régimen tributario *',
    'Tipo de cliente *',
    'Moneda principal *',
    'Tarifa de precios *',
    'Actividad económica CIUU',
    'ID EFFI: Forma de pago',
    'ID EFFI: Retención',
    'Permitir venta *',
    'Descuento %',
    'Cupo de crédito',
    'ID EFFI: Tipo de marketing',
    'ID EFFI: Sucursal',
    'ID EFFI: Ruta logística',
    'Email usuario responsable',
    'ID EFFI: Vendedor',
    'Restringir sucursal',
    'Restringir ruta logística',
    'Observación',
]

# ─── Cargar lookup tables ───────────────────────────────────────────────────────

def cargar_tarifas(conn_effi):
    """Devuelve dict {nombre_tarifa_lower: id_effi}."""
    cur = conn_effi.cursor()
    cur.execute("SELECT id, nombre FROM zeffi_tarifas_precios WHERE id IS NOT NULL AND nombre IS NOT NULL")
    mapa = {nombre.strip().lower(): str(id_) for id_, nombre in cur.fetchall()}
    cur.close()
    return mapa


def cargar_tipos_marketing(conn_effi):
    """Devuelve dict {nombre_marketing_lower: id_effi}."""
    cur = conn_effi.cursor()
    cur.execute("SELECT id, tipo_de_marketing FROM zeffi_tipos_marketing WHERE id IS NOT NULL AND tipo_de_marketing IS NOT NULL")
    mapa = {nombre.strip().lower(): str(id_) for id_, nombre in cur.fetchall()}
    cur.close()
    return mapa


def cargar_ciudades_id(conn_espo):
    """Devuelve dict {nombre_lower: id_effi} para buscar por nombre de ciudad."""
    cur = conn_espo.cursor()
    cur.execute("SELECT name, id_effi FROM ciudad WHERE deleted=0 AND id_effi IS NOT NULL")
    mapa = {(nombre or '').strip().lower(): id_effi for nombre, id_effi in cur.fetchall()}
    cur.close()
    return mapa


def cargar_vendedores(conn_effi):
    """Devuelve dict {nombre_completo_lower: codigo_effi}."""
    cur = conn_effi.cursor()
    cur.execute("SELECT codigo, nombre_completo FROM zeffi_empleados WHERE codigo IS NOT NULL AND nombre_completo IS NOT NULL")
    mapa = {nombre.strip().lower(): str(codigo) for codigo, nombre in cur.fetchall()}
    cur.close()
    return mapa


# ─── Leer contactos CRM ────────────────────────────────────────────────────────

def leer_contactos_crm(conn_espo, todos=False):
    """Lee contactos con fuente='CRM' y (si no --todos) enviado_a_effi=0."""
    cur = conn_espo.cursor(dictionary=True, buffered=True)

    filtro_enviado = '' if todos else 'AND (c.enviado_a_effi = 0 OR c.enviado_a_effi IS NULL)'

    cur.execute(f"""
        SELECT
            c.id,
            c.first_name,
            c.last_name,
            c.address_street,
            c.address_city,
            c.address_state,
            c.address_country,
            c.description,
            c.numero_identificacion,
            c.tipo_identificacion,
            c.tipo_persona,
            c.tipo_cliente,
            c.tipo_de_marketing,
            c.tarifa_precios,
            c.forma_pago,
            c.vendedor_effi,
            c.ciudad_nombre,
            -- email primario
            (SELECT ea.name FROM entity_email_address eea
             JOIN email_address ea ON ea.id = eea.email_address_id
             WHERE eea.entity_id = c.id AND eea.entity_type = 'Contact'
               AND eea.deleted = 0 AND eea.`primary` = 1 LIMIT 1) AS email,
            -- teléfono primario
            (SELECT pn.name FROM entity_phone_number epn
             JOIN phone_number pn ON pn.id = epn.phone_number_id
             WHERE epn.entity_id = c.id AND epn.entity_type = 'Contact'
               AND epn.deleted = 0 AND epn.`primary` = 1 LIMIT 1) AS telefono_principal,
            -- teléfono secundario
            (SELECT pn.name FROM entity_phone_number epn
             JOIN phone_number pn ON pn.id = epn.phone_number_id
             WHERE epn.entity_id = c.id AND epn.entity_type = 'Contact'
               AND epn.deleted = 0 AND epn.`primary` = 0 LIMIT 1) AS telefono_secundario
        FROM contact c
        WHERE c.deleted = 0
          AND c.fuente = 'CRM'
          AND (c.numero_identificacion IS NOT NULL AND c.numero_identificacion != '')
          {filtro_enviado}
        ORDER BY c.created_at
    """)
    contactos = cur.fetchall()
    cur.close()
    return contactos


# ─── Construir fila ────────────────────────────────────────────────────────────

def construir_fila(c, tarifas_map, marketing_map, ciudades_id_map, vendedores_map):
    """Convierte un contacto EspoCRM a una fila de la plantilla Effi."""

    nombre_completo = ' '.join(filter(None, [c.get('first_name'), c.get('last_name')])).strip()

    tipo_id_effi = TIPO_ID_MAP.get(c.get('tipo_identificacion') or '', None)

    tipo_persona_num = TIPO_PERSONA_MAP.get(c.get('tipo_persona') or '', 1)

    regimen = REGIMEN_JURIDICA if tipo_persona_num == 2 else REGIMEN_NATURAL

    tarifa_id = None
    if c.get('tarifa_precios'):
        tarifa_id = tarifas_map.get(c['tarifa_precios'].strip().lower())

    marketing_id = None
    if c.get('tipo_de_marketing'):
        marketing_id = marketing_map.get(c['tipo_de_marketing'].strip().lower())

    ciudad_effi_id = None
    if c.get('ciudad_nombre'):
        ciudad_effi_id = ciudades_id_map.get(c['ciudad_nombre'].strip().lower())

    # Tipo de cliente → ID numérico Effi
    tipo_cliente = TIPO_CLIENTE_MAP.get(c.get('tipo_cliente') or '', None)

    # Vendedor → código numérico Effi (lookup por nombre completo)
    vendedor_id = None
    if c.get('vendedor_effi'):
        vendedor_id = vendedores_map.get(c['vendedor_effi'].strip().lower())

    fila = [
        tipo_id_effi,                           # 0: ID EFFI: Tipo de documento *
        c.get('numero_identificacion'),          # 1: Número de documento *
        nombre_completo or None,                 # 2: Nombre | Razón social *
        c.get('email'),                          # 3: Email
        None,                                    # 4: Página web
        ciudad_effi_id,                          # 5: ID EFFI: Ciudad
        None,                                    # 6: Código DANE Ciudad
        c.get('address_street'),                 # 7: Dirección
        c.get('telefono_principal'),             # 8: Teléfono 1
        None,                                    # 9: Referencia teléfono 1
        c.get('telefono_secundario'),            # 10: Teléfono 2
        None,                                    # 11: Referencia teléfono 2
        None,                                    # 12: Celular (ya en teléfono 1)
        None,                                    # 13: WhatsApp
        None,                                    # 14: Fecha de nacimiento
        None,                                    # 15: Género
        None,                                    # 16: Estado cívil
        tipo_persona_num,                        # 17: Tipo de persona *
        regimen,                                 # 18: Régimen tributario *
        tipo_cliente,                            # 19: Tipo de cliente *
        'COP',                                   # 20: Moneda principal *
        tarifa_id,                               # 21: Tarifa de precios *
        None,                                    # 22: Actividad económica CIUU
        None,                                    # 23: ID EFFI: Forma de pago
        None,                                    # 24: ID EFFI: Retención
        1,                                       # 25: Permitir venta *
        None,                                    # 26: Descuento %
        None,                                    # 27: Cupo de crédito
        marketing_id,                            # 28: ID EFFI: Tipo de marketing
        1,                                       # 29: ID EFFI: Sucursal
        None,                                    # 30: ID EFFI: Ruta logística
        EMAIL_RESPONSABLE,                       # 31: Email usuario responsable
        vendedor_id,                             # 32: ID EFFI: Vendedor
        None,                                    # 33: Restringir sucursal
        None,                                    # 34: Restringir ruta logística
        c.get('description'),                    # 35: Observación
    ]
    return fila


# ─── Generar XLSX ──────────────────────────────────────────────────────────────

def generar_xlsx(contactos, tarifas_map, marketing_map, ciudades_id_map, vendedores_map, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Clientes'

    # Encabezados
    ws.append(COLUMNAS)

    # Datos
    for c in contactos:
        fila = construir_fila(c, tarifas_map, marketing_map, ciudades_id_map, vendedores_map)
        ws.append(fila)

    wb.save(str(output_path))
    return len(contactos)


# ─── Marcar como enviados ──────────────────────────────────────────────────────

def marcar_enviados(conn_espo, ids):
    if not ids:
        return
    cur = conn_espo.cursor()
    placeholders = ', '.join(['%s'] * len(ids))
    cur.execute(
        f"UPDATE contact SET enviado_a_effi = 1 WHERE id IN ({placeholders})",
        ids
    )
    conn_espo.commit()
    cur.close()


# ─── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--todos', action='store_true',
                        help='Incluir contactos ya enviados (re-exportar todos los CRM)')
    parser.add_argument('--marcar', action='store_true', default=True,
                        help='Marcar contactos como enviados tras generar (default: True)')
    parser.add_argument('--no-marcar', dest='marcar', action='store_false',
                        help='No marcar contactos como enviados')
    args = parser.parse_args()

    conn_espo = mysql.connector.connect(**DB_ESPO, autocommit=False)
    conn_effi = mysql.connector.connect(**DB_EFFI)

    try:
        # Cargar lookups
        tarifas_map     = cargar_tarifas(conn_effi)
        marketing_map   = cargar_tipos_marketing(conn_effi)
        ciudades_id_map = cargar_ciudades_id(conn_espo)
        vendedores_map  = cargar_vendedores(conn_effi)

        # Leer contactos
        contactos = leer_contactos_crm(conn_espo, todos=args.todos)

        if not contactos:
            print('✅ generar_plantilla_import_effi — sin contactos CRM pendientes')
            return

        # Generar XLSX
        hoy = date.today().strftime('%Y-%m-%d')
        output_path = OUTPUT_DIR / f'import_clientes_effi_{hoy}.xlsx'
        total = generar_xlsx(contactos, tarifas_map, marketing_map, ciudades_id_map, vendedores_map, output_path)

        print(f'✅ generar_plantilla_import_effi — {total} contactos → {output_path}')

        # Marcar como enviados
        if args.marcar and not args.todos:
            ids = [c['id'] for c in contactos]
            marcar_enviados(conn_espo, ids)
            print(f'   {total} contactos marcados como enviado_a_effi=1')
        elif args.todos:
            print('   (modo --todos: no se marca enviado_a_effi)')

    finally:
        conn_espo.close()
        conn_effi.close()


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        print(f'❌ generar_plantilla_import_effi — ERROR: {e}', file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)
